# -*- coding: utf-8 -*-
"""Сборка книги «Консигнация CAT» для чистого Excel: данные — Power Query из выгрузок SAP,
расчёт — вычисляемые столбцы таблицы «Позиции», которую заполняет запрос.

    python consignment/excel/make_pq.py            # книга с запросами и предзаполненной таблицей
    python consignment/excel/make_pq.py --no-pq --out /tmp/t.xlsx   # без запросов — для проверки формул

Книгу потом можно вести целиком в Excel: Python нужен только чтобы собрать её заново
после правок формул. Код запросов — section_template.m (пишется и рядом с книгой как *.Section1.m).
"""
import argparse, base64, datetime, io, json, os, re, struct, uuid, zipfile
from urllib.parse import quote
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
ap = argparse.ArgumentParser()
ap.add_argument("--data", default=os.path.join(os.path.dirname(HERE), "cons_data.json"), help="данные для предзаполнения таблицы")
ap.add_argument("--m", default=os.path.join(HERE, "section_template.m"))
ap.add_argument("--out", default=os.path.join(HERE, "Консигнация_CAT_Excel_PQ.xlsx"))
ap.add_argument("--limit", type=int, default=0)
ap.add_argument("--no-pq", action="store_true", help="без встроенных запросов (для проверки формул в LibreOffice)")
A = ap.parse_args()
D = json.load(open(A.data))

F = "Arial"
f_base = Font(name=F, size=10); f_in = Font(name=F, size=10, color="0000FF"); f_link = Font(name=F, size=10, color="008000")
f_bold = Font(name=F, size=10, bold=True); f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=14, bold=True); f_note = Font(name=F, size=9, italic=True, color="555555")
fill_in = PatternFill("solid", fgColor="FFFF00"); fill_head = PatternFill("solid", fgColor="2A3138")
fill_sub = PatternFill("solid", fgColor="E6F7F0")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top"); center = Alignment(horizontal="center", vertical="center", wrap_text=True)
RUB = '#,##0;(#,##0);"-"'; MLN = '#,##0.0;(#,##0.0);"-"'; QTY = '#,##0.##;(#,##0.##);"-"'
INT = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; PCT0 = '0%;(0%);"-"'
MON = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
T = "Позиции"

# ---------------- исходные столбцы (их отдаёт запрос «Позиции») ----------------
# (внутреннее имя в M, заголовок, тип M, формат, ключ данных)
SRC = [("Площадка", "Площадка", "text", "@", "site"), ("№", "№", "int", INT, "_n"),
       ("Номер Cat", "Номер Cat", "text", "@", "no"), ("ЕКМТР", "Код ЕК МТР", "text", "@", "ek"),
       ("Наименование", "Наименование", "text", "@", "name"), ("Модели", "Модели", "text", "@", "models"),
       ("Часть", "Часть (поле / цех)", "text", "@", "part"), ("Признак", "Признак CAPEX / OPEX", "text", "@", "capex"),
       ("Причина", "Причина инвестиций SAP (план)", "text", "@", "rs"), ("Источник", "Источник спроса", "text", "@", "src"),
       ("Виды работ", "Виды работ", "text", "@", "work"),
       ("План шт", "План, шт", "num", QTY, "q27"), ("План поле", "в т.ч. поле, шт", "num", QTY, "q27_pole"),
       ("План цех", "в т.ч. цех, шт", "num", QTY, "q27_ceh"), ("План CAPEX", "в т.ч. CAPEX, шт", "num", QTY, "q27_capex"),
       ("План ВНЕ", "в т.ч. внеплановые виды работ в плане, шт", "num", QTY, "q27_vne"),
       ("Заказов план", "Заказов в плане", "num", INT, "picks"), ("Макс отбор план", "Макс. разовый отбор в плане, шт", "num", QTY, "maxpick")] + \
      [("П%d" % (i + 1), "План %s, шт" % m, "num", QTY, ("m27", i)) for i, m in enumerate(MON)] + \
      [("Заказов -%d" % k, "Вне плана: заказов (год −%d)" % k, "num", INT, ("h_o", 3 - k)) for k in (3, 2, 1)] + \
      [("Шт -%d" % k, "Вне плана: шт (год −%d)" % k, "num", QTY, ("h_q", 3 - k)) for k in (3, 2, 1)] + \
      [("Макс отбор вне", "Вне плана: макс. разовый отбор, шт", "num", QTY, "h_maxpick"),
       ("Доля поле вне", "Вне плана: доля поле", "num", PCT0, "h_pole_share"),
       ("Доля CAPEX вне", "Вне плана: доля CAPEX", "num", PCT0, "h_capex_share"),
       ("Цена", "Цена, руб/шт", "num", RUB, "price"),
       ("Ост БЕ", "Остаток БЕ площадки, шт", "num", QTY, "st_be"),
       ("Ост Развитие", "Остаток АО «Развитие» без подрядчиков, шт", "num", QTY, "st_razv"),
       ("Ост замены", "Остаток по заменам Cat, шт", "num", QTY, "st_rep"),
       ("Номера-замены", "Номера-замены", "text", "@", "rep"),
       ("Ост Подрядчик", "Справочно: на складах подрядчиков, шт", "num", QTY, "st_contr"),
       ("Ост Консигнация", "Справочно: действующая консигнация, шт", "num", QTY, "st_cons"),
       ("Ост АТЗ", "Справочно: АТЗ, шт", "num", QTY, "st_atz"),
       ("Ост Другой завод", "Справочно: остаток на других заводах региона, шт", "num", QTY, "st_other"),
       ("Пост до года", "Поставки: до года плана и просроченные, шт", "num", QTY, "del_open"),
       ("Пост просрочено", "в т.ч. просрочено на дату выгрузки, шт", "num", QTY, "del_late")] + \
      [("Пост м%d" % (i + 1), "Поставка %s, шт" % m, "num", QTY, ("del_m", i)) for i, m in enumerate(MON)] + \
      [("Пост после", "Поставки после года плана (справочно), шт", "num", QTY, "del_after"),
       ("Пост неподтв", "Неподтверждённые: заявки и заказы на согласовании, шт", "num", QTY, "del_pend"),
       ("Поставщики", "Поставщики по заказам", "text", "@", "del_sup"),
       ("Заказов закупки", "Заказов закупки", "int", INT, "del_docs"),
       ("Пост другой завод", "Справочно: заказы на другие заводы региона, шт", "num", QTY, "del_other")]
H = {s[1]: s for s in SRC}

# ---------------- параметры площадок ----------------
PARAMS = [  # (заголовок в таблице параметров, имя столбца-помощника в «Позициях», 1100, 1300, формат, пояснение)
 ("Срок пополнения, дн", "L, дн", 90, 90, INT, "От заказа пополнения до поступления на склад консигнации."),
 ("Период пересмотра, дн", "R, дн", 7, 7, INT, "Как часто подрядчик сверяет остаток с Min."),
 ("Уровень сервиса", "SL", 0.95, 0.95, PCT, "Вероятность закрыть внеплановые заказы за срок пополнения."),
 ("Коэффициент выполнения плана", "k плана", 1.0, 1.0, "0.00", "1,00 — строго по плану ТОРО."),
 ("Учитывать спрос вне плана (1/0)", "Вне плана вкл.", 1, 1, "0", "Расход на внеплановых работах за три прошлых года."),
 ("Порог вероятности", "Порог вероятности", 0.5, 0.5, PCT0, "Внеплановый спрос в расчёте, если вероятность использования не ниже порога."),
 ("Длительность истории, лет", "Лет истории", 2.667, 2.667, "0.00", "Три года до года плана; последний обычно неполный (янв–авг = 2,67)."),
 ("Коэффициент к истории", "k истории", 1.0, 1.0, "0.00", "Масштаб внепланового спроса (изменение парка и наработки)."),
 ("Учитывать остаток «Развития» (1/0)", "Развитие вкл.", 1, 1, "0", "Сначала вырабатываем свой остаток на заводах площадки."),
 ("Учитывать замены Cat (1/0)", "Замены вкл.", 1, 1, "0", "Остаток под взаимозаменяемыми номерами Caterpillar."),
 ("Лимит стоимости, руб (0 — без лимита)", "Лимит, руб", 0, 0, RUB, "Лимит на запас на площадке (макс.)."),
 ("Учитывать поставки по заказам (1/0)", "Заказы вкл.", 1, 1, "0", "Действующие заказы АО «Развитие» на завод площадки."),
 ("Учитывать неподтверждённые (1/0)", "Неподтв. вкл.", 0, 0, "0", "Заявки без заказа и заказы на согласовании."),
 ("Задержка поставок, мес", "Задержка, мес", 0, 0, "0", "Сдвиг всех поставок по заказам."),
]
PT = "ПараметрыПлощадок"

def q(name):  # экранирование имени столбца в структурной ссылке
    return re.sub(r"([\[\]#'])", r"'\1", name)
def r(name): return "%s[[#This Row],[%s]]" % (T, q(name))
def c(name): return "%s[[%s]]" % (T, q(name))
def rr(a, b): return "%s[[#This Row],[%s]:[%s]]" % (T, q(a), q(b))
KS = "{" + ",".join(str(i) for i in range(60)) + "}"
def pq(m, sl):
    return ("IF(%s<20,SUMPRODUCT((POISSON(%s,%s,TRUE)<%s)*1),"
            "CEILING(%s+NORMSINV(%s)*SQRT(%s)+(NORMSINV(%s)^2-1)/6-0.5,1))") % (m, KS, m, sl, m, sl, m, sl)

P = {p[1]: r(p[1]) for p in PARAMS}
L, R_, SL, K, SWV, THR, HY, KH, SWR, SWP, BUD, SWD, SWQ, DLY = (P[k] for k in
    ("L, дн", "R, дн", "SL", "k плана", "Вне плана вкл.", "Порог вероятности", "Лет истории", "k истории",
     "Развитие вкл.", "Замены вкл.", "Лимит, руб", "Заказы вкл.", "Неподтв. вкл.", "Задержка, мес"))
W = r("W, мес")

CALC = []  # (заголовок, формула без «=», формат)
for p in PARAMS:
    CALC.append((p[1], "INDEX(%s[[%s]],MATCH(%s,%s[[Площадка]],0))" % (PT, q(p[0]), r("Площадка"), PT), p[4]))
CALC.append(("W, мес", "MAX(1,MIN(12,ROUND(%s/30.4,0)))" % L, "0"))
ho = "SUM(%s)" % rr("Вне плана: заказов (год −3)", "Вне плана: заказов (год −1)")
hq = "SUM(%s)" % rr("Вне плана: шт (год −3)", "Вне плана: шт (год −1)")
lam, qbar, prob, inc = r("Вне плана: заказов в год"), r("Вне плана: шт на заказ"), r("Вероятность использования в году плана"), r("Вне плана в расчёт (1/0)")
dp, dv, dem, nord = r("Спрос по плану, шт/год"), r("Спрос вне плана в расчёт, шт/год"), r("Потребность расчётная, шт/год"), r("Заказов в год")
price = r("Цена, руб/шт")
CALC += [
 ("Вне плана: заказов в год", "%s/%s*%s" % (ho, HY, KH), "0.00"),
 ("Вне плана: шт на заказ", "IF(%s>0,%s/%s,0)" % (ho, hq, ho), QTY),
 ("Вероятность использования в году плана", "1-EXP(-%s)" % lam, PCT0),
 ("Вне плана в расчёт (1/0)", "IF(AND(%s=1,%s>0,%s>=%s),1,0)" % (SWV, lam, prob, THR), "0"),
 ("Вне плана: ожидаемо, шт/год (без порога)", "%s*%s" % (lam, qbar), QTY),
 ("Вне плана: ожидаемо, руб/год (без порога)", "%s*%s" % (r("Вне плана: ожидаемо, шт/год (без порога)"), price), RUB),
 ("Спрос по плану, шт/год", "%s*%s" % (r("План, шт"), K), QTY),
 ("Спрос вне плана в расчёт, шт/год", "%s*MAX(0,%s-%s*%s)" % (inc, r("Вне плана: ожидаемо, шт/год (без порога)"), r("в т.ч. внеплановые виды работ в плане, шт"), K), QTY),
 ("Потребность расчётная, шт/год", "%s+%s" % (dp, dv), QTY),
 ("Заказов в год", "%s*%s+IF(%s>0,%s,0)" % (r("Заказов в плане"), K, dv, lam), "0.0"),
]
for i, m in enumerate(MON):
    prev = "" if i == 0 else r("План нараст. %s" % MON[i - 1]) + "+"
    CALC.append(("План нараст. %s" % m, "%s%s*%s" % (prev, r("План %s, шт" % m), K), QTY))
terms = ["%s-IF(%d>%s,INDEX(%s,1,%d-%s),0)" % (r("План нараст. %s" % MON[j - 1]), j, W, rr("План нараст. янв", "План нараст. дек"), j, W) for j in range(1, 13)]
mplan = r("План: макс. потребность за срок пополнения, шт")
qe, mu, mur = r("Вне плана: шт на заказ в расчёте"), r("Вне плана: заказов за срок пополнения"), r("Вне плана: заказов за срок + пересмотр")
mn, mx = r("Min, шт"), r("Max, шт")
maxpick = "CEILING(MAX(%s,IF(%s>0,%s,0)),1)" % (r("Макс. разовый отбор в плане, шт"), dv, r("Вне плана: макс. разовый отбор, шт"))
CALC += [
 ("План: макс. потребность за срок пополнения, шт", "MAX(%s)" % ",".join(terms), QTY),
 ("Вне плана: шт на заказ в расчёте", "IF(AND(%s>0,%s>0),%s/%s,0)" % (dv, lam, dv, lam), QTY),
 ("Вне плана: заказов за срок пополнения", "IF(%s>0,%s*%s/365,0)" % (dv, lam, L), "0.00"),
 ("Вне плана: заказов за срок + пересмотр", "IF(%s>0,%s*(%s+%s)/365,0)" % (dv, lam, L, R_), "0.00"),
 ("Min, шт", "IF(%s<=0,0,CEILING(%s+IF(%s>0,%s*%s,0)-0.000001,1))" % (dem, mplan, mu, pq(mu, SL), qe), INT),
 ("Max, шт", "IF(%s<=0,0,MAX(CEILING(%s+%s*%s/365+IF(%s>0,%s*%s,0)-0.000001,1),%s+1,%s))" % (dem, mplan, dp, R_, mur, pq(mur, SL), qe, mn, maxpick), INT),
 ("Собственный остаток к выработке, шт", "%s+%s*%s+%s*%s" % (r("Остаток БЕ площадки, шт"), r("Остаток АО «Развитие» без подрядчиков, шт"), SWR, r("Остаток по заменам Cat, шт"), SWP), QTY),
]
own = r("Собственный остаток к выработке, шт")
for i, m in enumerate(MON):
    prev = "" if i == 0 else r("Поставки нараст. %s" % MON[i - 1]) + "+"
    CALC.append(("Поставки нараст. %s" % m, "%s%s" % (prev, r("Поставка %s, шт" % m)), QTY))
for i, m in enumerate(MON):
    k = i + 1
    CALC.append(("Своё предложение нараст. %s" % m,
        "%s+%s*(IF(%d>=%s,%s,0)+IF(%d-%s>=1,INDEX(%s,1,%d-%s),0))+%s*IF(%d>=MIN(12,%s+1+%s),%s,0)" % (
            own, SWD, k, DLY, r("Поставки: до года плана и просроченные, шт"), k, DLY,
            rr("Поставки нараст. янв", "Поставки нараст. дек"), k, DLY, SWQ, k, W, DLY,
            r("Неподтверждённые: заявки и заказы на согласовании, шт")), QTY))
CALC += [
 ("Поставки в расчёт, шт", "%s-%s" % (r("Своё предложение нараст. дек"), own), QTY),
 ("Своё предложение на год (остаток + поставки), шт", r("Своё предложение нараст. дек"), QTY),
]
for i, m in enumerate(MON):
    prev = "" if i == 0 else r("Нараст. %s" % MON[i - 1]) + "+"
    CALC.append(("Нараст. %s" % m, "%s%s*%s+%s/12" % (prev, r("План %s, шт" % m), K, dv), QTY))
entry = ["IF(%s>%s-%s,%d,13)" % (r("Нараст. %s" % m), r("Своё предложение нараст. %s" % m), mn, i + 1) for i, m in enumerate(MON)]
mon = r("Месяц ввода в консигнацию"); inn = r("В консигнации в году плана (1/0)")
edl, ss, onmax = r("Спрос за срок пополнения (среднее), шт"), r("Страховой запас, шт"), r("Макс. запас на площадке, шт")
val = r("Стоимость на площадке (макс.), руб"); prio = r("Приоритет: заказов на 1 млн руб"); rk = r("Ранг")
CALC += [
 ("Месяц ввода в консигнацию", "IF(%s<=0,13,MIN(%s))" % (dem, ",".join(entry)), '[<13]0;"после года"'),
 ("В консигнации в году плана (1/0)", "IF(AND(%s>0,%s<=12),1,0)" % (dem, mon), "0"),
 ("Спрос за срок пополнения (среднее), шт", "%s*%s/365" % (dem, L), QTY),
 ("Страховой запас, шт", "MAX(0,%s-%s)" % (mn, edl), QTY),
 ("Макс. запас на площадке, шт", "IF(%s<=0,0,MAX(CEILING(%s-%s-0.000001,1),%s,1))" % (dem, mx, edl, maxpick), INT),
 ("Стоимость по Max с учётом в пути, руб", "%s*%s*%s" % (inn, mx, price), RUB),
 ("Стоимость на площадке (макс.), руб", "%s*%s*%s" % (inn, onmax, price), RUB),
 ("Средний запас на площадке, руб", "%s*MIN(%s,%s+(%s-%s)/2)*%s" % (inn, onmax, ss, mx, mn, price), RUB),
 ("Выборка с консигнации за год, руб", "%s*MAX(0,%s-%s)*%s" % (inn, dem, r("Своё предложение на год (остаток + поставки), шт"), price), RUB),
 ("Доля поле", "IF(%s>0,(%s*%s+%s*%s)/%s,0)" % (dem, r("в т.ч. поле, шт"), K, dv, r("Вне плана: доля поле"), dem), PCT0),
 ("Доля CAPEX", "IF(%s>0,(%s*%s+%s*%s)/%s,0)" % (dem, r("в т.ч. CAPEX, шт"), K, dv, r("Вне плана: доля CAPEX"), dem), PCT0),
 ("Доля вне плана", "IF(%s>0,%s/%s,0)" % (dem, dv, dem), PCT0),
 ("Приоритет: заказов на 1 млн руб", "IF(%s>0,%s/(%s/1000000),0)" % (val, nord, val), "0.0"),
 ("Ранг", 'IF(%s>0,COUNTIFS(%s,%s,%s,">0",%s,">"&%s)+COUNTIFS(%s,%s,%s,">0",%s,%s,%s,"<"&%s)+1,0)' % (
     val, c("Площадка"), r("Площадка"), c("Стоимость на площадке (макс.), руб"), c("Приоритет: заказов на 1 млн руб"), prio,
     c("Площадка"), r("Площадка"), c("Стоимость на площадке (макс.), руб"), c("Приоритет: заказов на 1 млн руб"), prio, c("№"), r("№")), INT),
 ("Накопленная стоимость, руб", 'IF(%s>0,SUMIFS(%s,%s,%s,%s,"<="&%s,%s,">0"),0)' % (
     rk, c("Стоимость на площадке (макс.), руб"), c("Площадка"), r("Площадка"), c("Ранг"), rk, c("Ранг")), RUB),
 ("В лимите (1/0)", "IF(%s=0,0,IF(OR(%s=0,%s<=%s),1,0))" % (inn, BUD, r("Накопленная стоимость, руб"), BUD), "0"),
 ("Стоимость в лимите, руб", "%s*%s" % (r("В лимите (1/0)"), val), RUB),
 ("Заказов в лимите", "%s*%s" % (r("В лимите (1/0)"), nord), "0.0"),
]
HIDE = [h for h, _, _ in CALC if h.startswith(("План нараст.", "Поставки нараст.", "Своё предложение нараст.", "Нараст."))] + [p[1] for p in PARAMS]
KEY = ["Min, шт", "Max, шт", "Месяц ввода в консигнацию", "Макс. запас на площадке, шт", "Стоимость на площадке (макс.), руб", "В лимите (1/0)", "Вероятность использования в году плана"]
ALLCOLS = [s[1] for s in SRC] + [h for h, _, _ in CALC]
assert len(ALLCOLS) == len(set(ALLCOLS)), [x for x in ALLCOLS if ALLCOLS.count(x) > 1]

# ---------------- книга ----------------
wb = Workbook()
Pz = wb.active; Pz.title = "Параметры"
Pz["A1"] = "Внешняя консигнация подрядчика CAT — параметры"; Pz["A1"].font = f_title
Pz["A2"] = "Жёлтые ячейки — вводимые. После смены папки или года: Данные → Обновить все. Остальные параметры пересчитываются сразу."; Pz["A2"].font = f_note
Pz["A4"] = "Загрузка данных (Power Query)"; Pz["A4"].font = f_bold
LOAD = [("Папка с выгрузками", "C:\\Консигнация\\Выгрузки", "ПапкаДанных", "@",
         "M06_<площадка>_<год>.xlsx (PM-06), Остатки*.xlsx (MM-M03), Закупка*.xlsx, ЕКМТР*.xlsx, Замены*.xlsx"),
        ("Год плана", 2027, "ГодПлана", "0", "План — M06_<площадка>_<год плана>, история вне плана — три года до него"),
        ("Дата выгрузки закупки", datetime.date(2026, 9, 11), "ДатаВыгрузкиЗакупки", "DD.MM.YYYY", "Поставки с датой раньше неё — просроченные")]
for i, (lab, v, nm, fmt, note) in enumerate(LOAD):
    rrow = 5 + i
    Pz.cell(rrow, 1, lab).font = f_base
    cc = Pz.cell(rrow, 2, v); cc.font = f_in; cc.fill = fill_in; cc.border = box; cc.number_format = fmt
    Pz.cell(rrow, 3, note).font = f_note
    wb.defined_names[nm] = DefinedName(nm, attr_text="'Параметры'!$B$%d" % rrow)
Pz["A10"] = "Параметры расчёта по площадкам"; Pz["A10"].font = f_bold
PR0 = 11
Pz.cell(PR0, 1, "Площадка")
for j, p in enumerate(PARAMS): Pz.cell(PR0, 2 + j, p[0])
for i, site in enumerate(("1100", "1300")):
    cc = Pz.cell(PR0 + 1 + i, 1, site); cc.font = f_bold; cc.number_format = "@"
    for j, p in enumerate(PARAMS):
        cc = Pz.cell(PR0 + 1 + i, 2 + j, p[2 + i]); cc.font = f_in; cc.fill = fill_in; cc.number_format = p[4]; cc.border = box
for j in range(len(PARAMS) + 1):
    h = Pz.cell(PR0, 1 + j); h.font = f_head; h.fill = fill_head; h.alignment = center
tp = Table(displayName=PT, ref="A%d:%s%d" % (PR0, CL(len(PARAMS) + 1), PR0 + 2))
tp.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False); Pz.add_table(tp)
Pz.row_dimensions[PR0].height = 58
Pz.column_dimensions["A"].width = 26; Pz.column_dimensions["B"].width = 22; Pz.column_dimensions["C"].width = 22
for j in range(3, len(PARAMS) + 2): Pz.column_dimensions[CL(j + 1)].width = 15
Pz.cell(PR0 + 3, 1, "1100 — Красноярск (793D), 1300 — Алдан (773E, 777G, 395)").font = f_note
Pz.cell(PR0 + 5, 1, "Пояснения").font = f_bold
for i, p in enumerate(PARAMS):
    Pz.cell(PR0 + 6 + i, 1, p[0]).font = f_base; Pz.cell(PR0 + 6 + i, 2, p[5]).font = f_note
for rng_, f1 in (("B%d:B%d" % (PR0 + 1, PR0 + 2), '"30,45,60,90,120,150,180"'),):
    dv_ = DataValidation(type="list", formula1=f1); Pz.add_data_validation(dv_); dv_.add(rng_)

# настройки: техника, заводы, исключения
S = wb.create_sheet("Настройки")
S["A1"] = "Настройки загрузки — Power Query читает эти таблицы"; S["A1"].font = f_title
S["A2"] = "Можно добавлять строки: новую технику, заводы, правила исключения. После правки — Данные → Обновить все."; S["A2"].font = f_note
def cfg_table(col0, name, head, rows, widths):
    S.cell(4, col0, {"Техника": "Техника площадок (ЕО начинается с…)", "Заводы": "Заводы площадок",
                     "Исключения": "Исключения из номенклатуры (наименование начинается с…)"}[name]).font = f_bold
    for j, h in enumerate(head):
        cc = S.cell(5, col0 + j, h); cc.font = f_head; cc.fill = fill_head; cc.alignment = center
        S.column_dimensions[CL(col0 + j)].width = widths[j]
    for i, rw in enumerate(rows):
        for j, v in enumerate(rw):
            cc = S.cell(6 + i, col0 + j, v); cc.font = f_in; cc.number_format = "@"
    t = Table(displayName=name, ref="%s5:%s%d" % (CL(col0), CL(col0 + len(head) - 1), 5 + len(rows)))
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True); S.add_table(t)
cfg_table(1, "Техника", ["Площадка", "Начало наименования ЕО", "Модель"], [
    ("1100", "Самосвал карьерный 793D", "793D"), ("1100", "Кузов 793D", "793D"), ("1100", "ДВС 3516", "793D"),
    ("1300", "Самосвал карьерный 773E", "773E"), ("1300", "Самосвал карьерный 777G", "777G"),
    ("1300", "Экскаватор гидравл. 395", "395"), ("1300", "Экскаватор гидравлический 395", "395"),
    ("1300", "ДВС C27", "773E"), ("1300", "ДВС C32", "777G")], [10, 32, 10])
cfg_table(5, "Заводы", ["Площадка", "Код завода", "Роль"], [
    ("1100", "1100", "Остатки"), ("1100", "7101", "Остатки и закупка"), ("1100", "7106", "Справочно"),
    ("1300", "1300", "Остатки"), ("1300", "7103", "Остатки и закупка")], [10, 12, 20])
S.cell(12, 5, "Роли: «Остатки», «Закупка», «Остатки и закупка» — считаем своим; «Справочно» — только показываем (перемещение).").font = f_note
EXC = [("масло", "масла и смазки"), ("смазка", "масла и смазки"), ("жидкость", "масла и смазки"), ("антифриз", "масла и смазки"),
       ("охлаждающая жидк", "масла и смазки"), ("шина", "шины", "шина заземл"), ("покрышка", "шины"), ("камера шин", "шины"),
       ("коронка", "GET"), ("палец коронки", "GET"), ("фиксатор коронки", "GET"), ("зуб ", "GET"), ("зуб.", "GET"), ("зуб,", "GET"),
       ("нож ", "GET"), ("нож.", "GET"), ("нож,", "GET"), ("резак", "GET"), ("накладка ковша", "GET"), ("протектор ковша", "GET"),
       ("угол отвала", "GET"), ("кромка", "GET"),
       ("двигатель", "агрегат"), ("дв. сб", "агрегат"), ("дв.сб", "агрегат"), ("кпп", "агрегат"), ("пер.конечн", "агрегат"),
       ("передача конечн", "агрегат"), ("диф. сб", "агрегат"), ("диф.сб", "агрегат"), ("дифференциал", "агрегат"),
       ("гидр/трансф", "агрегат"), ("гидротрансформ", "агрегат"), ("группа колеса", "агрегат"), ("ступица сб", "агрегат"),
       ("подвеска", "агрегат"), ("круг поворотный", "агрегат"), ("привод поворота", "агрегат"), ("редуктор поворот", "агрегат"),
       ("цилиндр подъемника", "агрегат"), ("цилиндр стрелы", "агрегат"), ("цилиндр ковша", "агрегат"), ("гидроцилиндр", "агрегат"),
       ("насос гл.", "агрегат"), ("главная передача", "агрегат"), ("мост ", "агрегат")]
cfg_table(9, "Исключения", ["Начало наименования", "Причина", "Кроме"], [(e[0], e[1], e[2] if len(e) > 2 else "") for e in EXC], [24, 18, 16])

# ---------------- лист «Позиции»: таблица запроса + вычисляемые столбцы ----------------
Z = wb.create_sheet("Позиции")
rows = []
for site in ("1100", "1300"):
    rs_ = sorted([x for x in D["rows"] if x["site"] == site], key=lambda x: -((x["q27"] + sum(x["h_q"]) / 2.667) * x["price"]))
    if A.limit: rs_ = rs_[:A.limit]
    for i, x in enumerate(rs_):
        x = dict(x); x["_n"] = i + 1; rows.append(x)
N = len(rows); LASTROW = 1 + N; NC = len(ALLCOLS)
for j, h in enumerate(ALLCOLS):
    cc = Z.cell(1, j + 1, h); cc.font = f_head; cc.fill = fill_head; cc.alignment = center
Z.row_dimensions[1].height = 70
for i, x in enumerate(rows):
    rrow = i + 2
    for j, s in enumerate(SRC):
        key = s[4]
        v = x[key[0]][key[1]] if isinstance(key, tuple) else x[key]
        cc = Z.cell(rrow, j + 1, v); cc.font = f_in; cc.number_format = s[3]
    for j, (h, fml, fmt) in enumerate(CALC):
        cc = Z.cell(rrow, len(SRC) + j + 1, "=" + fml); cc.font = f_base; cc.number_format = fmt
        if h in KEY: cc.fill = fill_sub
for j, h in enumerate(ALLCOLS):
    Z.column_dimensions[CL(j + 1)].width = 34 if h in ("Наименование",) else 11
    if h in HIDE: Z.column_dimensions[CL(j + 1)].hidden = True
Z.freeze_panes = "F2"
tz = Table(displayName=T, ref="A1:%s%d" % (CL(NC), LASTROW))
tz.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); Z.add_table(tz)
for h, t in (("Вероятность использования в году плана", "1 − e^(−заказов вне плана в год): вероятность хотя бы одного внепланового использования."),
             ("Min, шт", "План — детерминированно (максимум плана за W месяцев подряд), вне плана — квантиль Пуассона числа заказов × шт на заказ."),
             ("Макс. запас на площадке, шт", "Max минус средний спрос за срок пополнения (он в пути), не меньше разового отбора. От него — стоимость, лимит, приоритет."),
             ("Месяц ввода в консигнацию", "Первый месяц, когда нарастающая потребность превышает своё предложение (остаток + поставки) минус Min.")):
    Z.cell(1, ALLCOLS.index(h) + 1).comment = Comment(t, "модель")

# ---------------- сводка ----------------
Sv = wb.create_sheet("Сводка", 1)
Sv["A1"] = "Сводка: консигнация подрядчика по площадкам"; Sv["A1"].font = f_title
Sv["A2"] = "Все значения — формулы от таблицы «Позиции» и параметров."; Sv["A2"].font = f_note
for j, h in enumerate(["Показатель", "1100", "1300", "Итого"]):
    cc = Sv.cell(4, j + 1, h); cc.font = f_head; cc.fill = fill_head; cc.alignment = center; cc.number_format = "@"
Sv.cell(5, 2, "Красноярск").font = f_note; Sv.cell(5, 3, "Алдан").font = f_note
Sv.column_dimensions["A"].width = 64
for col in "BCD": Sv.column_dimensions[col].width = 16
def par(s, name): return "INDEX(%s[[%s]],MATCH(%s,%s[[Площадка]],0))" % (PT, q(name), s, PT)
def si(s): return "%s,%s" % (c("Площадка"), s)
def sp(s, *cols): return "SUMPRODUCT((%s=%s)%s)" % (c("Площадка"), s, "".join("*" + x for x in cols))
M = [
 ("Срок пополнения, дней", lambda s: "=" + par(s, "Срок пополнения, дн"), INT, None),
 ("Порог вероятности вне плана", lambda s: "=" + par(s, "Порог вероятности"), PCT0, None),
 ("Позиций всего (план и/или вне плана)", lambda s: "=COUNTIFS(%s)" % si(s), INT, 1),
 ("   в т.ч. в плане", lambda s: '=COUNTIFS(%s,%s,">0")' % (si(s), c("План, шт")), INT, 1),
 ("   в т.ч. только вне плана", lambda s: '=COUNTIFS(%s,%s,"Вне плана")' % (si(s), c("Источник спроса")), INT, 1),
 ("Позиций со спросом вне плана в расчёте", lambda s: "=SUMIFS(%s,%s)" % (c("Вне плана в расчёт (1/0)"), si(s)), INT, 1),
 ("Позиций в консигнации", lambda s: "=SUMIFS(%s,%s)" % (c("В консигнации в году плана (1/0)"), si(s)), INT, 1),
 ("Позиций, где своего остатка и поставок хватает на год", lambda s: '=COUNTIFS(%s,%s,13,%s,">0")' % (si(s), c("Месяц ввода в консигнацию"), c("Потребность расчётная, шт/год")), INT, 1),
 ("Позиций в лимите", lambda s: "=SUMIFS(%s,%s)" % (c("В лимите (1/0)"), si(s)), INT, 1),
 (None,),
 ("Лимит, млн руб (0 — без лимита)", lambda s: "=%s/1000000" % par(s, "Лимит стоимости, руб (0 — без лимита)"), MLN, 1),
 ("Запас на площадке, макс., без лимита, млн руб", lambda s: "=SUMIFS(%s,%s)/1000000" % (c("Стоимость на площадке (макс.), руб"), si(s)), MLN, 1),
 ("Запас на площадке, макс., в лимите, млн руб", lambda s: "=SUMIFS(%s,%s)/1000000" % (c("Стоимость в лимите, руб"), si(s)), MLN, 1),
 ("   в т.ч. поле", lambda s: "=%s/1000000" % sp(s, c("Стоимость в лимите, руб"), c("Доля поле")), MLN, 1),
 ("   в т.ч. цех (КР агрегатов)", lambda s: "=%s/1000000" % sp(s, c("Стоимость в лимите, руб"), "(1-%s)" % c("Доля поле")), MLN, 1),
 ("   в т.ч. CAPEX", lambda s: "=%s/1000000" % sp(s, c("Стоимость в лимите, руб"), c("Доля CAPEX")), MLN, 1),
 ("   в т.ч. OPEX", lambda s: "=%s/1000000" % sp(s, c("Стоимость в лимите, руб"), "(1-%s)" % c("Доля CAPEX")), MLN, 1),
 ("   в т.ч. под спрос вне плана", lambda s: "=%s/1000000" % sp(s, c("Стоимость в лимите, руб"), c("Доля вне плана")), MLN, 1),
 ("Средний запас на площадке в лимите, млн руб", lambda s: "=%s/1000000" % sp(s, c("Средний запас на площадке, руб"), c("В лимите (1/0)")), MLN, 1),
 ("Справочно: уровень Max с учётом запаса в пути, млн руб", lambda s: "=%s/1000000" % sp(s, c("Стоимость по Max с учётом в пути, руб"), c("В лимите (1/0)")), MLN, 1),
 ("Выборка с консигнации за год в лимите, млн руб", lambda s: "=%s/1000000" % sp(s, c("Выборка с консигнации за год, руб"), c("В лимите (1/0)")), MLN, 1),
 (None,),
 ("Заказов в год (план + вне плана в расчёте)", lambda s: "=SUMIFS(%s,%s)" % (c("Заказов в год"), si(s)), INT, 1),
 ("Заказов в год по позициям в лимите", lambda s: "=SUMIFS(%s,%s)" % (c("Заказов в лимите"), si(s)), INT, 1),
 ("Доля заказов, закрываемых консигнацией", "ratio", PCT, "ratio"),
 (None,),
 ("Поставки по закупке в расчёте, млн руб", lambda s: "=%s/1000000" % sp(s, c("Поставки в расчёт, шт"), c("Цена, руб/шт")), MLN, 1),
 ("   в т.ч. просрочено на дату выгрузки (справочно), млн руб", lambda s: "=%s/1000000" % sp(s, c("в т.ч. просрочено на дату выгрузки, шт"), c("Цена, руб/шт")), MLN, 1),
 ("Справочно: внеплановый спрос по истории, млн руб/год (без порога)", lambda s: "=SUMIFS(%s,%s)/1000000" % (c("Вне плана: ожидаемо, руб/год (без порога)"), si(s)), MLN, 1),
 ("Справочно: остаток на других заводах региона, млн руб", lambda s: "=%s/1000000" % sp(s, c("Справочно: остаток на других заводах региона, шт"), c("Цена, руб/шт")), MLN, 1),
 ("Справочно: заказы на другие заводы региона, млн руб", lambda s: "=%s/1000000" % sp(s, c("Справочно: заказы на другие заводы региона, шт"), c("Цена, руб/шт")), MLN, 1),
 ("Справочно: остаток на складах подрядчиков, млн руб", lambda s: "=%s/1000000" % sp(s, c("Справочно: на складах подрядчиков, шт"), c("Цена, руб/шт")), MLN, 1),
]
rrow = 6; MR = {}
for m in M:
    if m[0] is None: rrow += 1; continue
    lab, fn, fmt, agg = m
    Sv.cell(rrow, 1, lab).font = f_base if lab.startswith("   ") else f_bold
    for j, col in enumerate("BC"):
        s = "%s$4" % col
        v = ("=IF(%s%d>0,%s%d/%s%d,0)" % (col, MR["Заказов в год (план + вне плана в расчёте)"], col, MR["Заказов в год по позициям в лимите"],
                                        col, MR["Заказов в год (план + вне плана в расчёте)"])) if fn == "ratio" else fn(s)
        cc = Sv.cell(rrow, 2 + j, v); cc.number_format = fmt; cc.border = box; cc.font = f_base
    if agg == 1: cc = Sv.cell(rrow, 4, "=B%d+C%d" % (rrow, rrow))
    elif agg == "ratio":
        a, b = MR["Заказов в год (план + вне плана в расчёте)"], MR["Заказов в год по позициям в лимите"]
        cc = Sv.cell(rrow, 4, "=IF(D%d>0,D%d/D%d,0)" % (a, b, a))
    else: cc = None
    if cc is not None: cc.number_format = fmt; cc.font = f_bold; cc.fill = fill_sub; cc.border = box
    MR[lab] = rrow; rrow += 1
rrow += 1
Sv.cell(rrow, 1, "Ввод позиций в консигнацию по месяцам (в лимите), запас на площадке макс., млн руб").font = f_bold; rrow += 1
for i, m in enumerate(["Площадка"] + MON):
    cc = Sv.cell(rrow, 1 + i, m); cc.font = f_head; cc.fill = fill_head; cc.alignment = center
for s, lab in (("1100", "1100 Красноярск"), ("1300", "1300 Алдан")):
    rrow += 1; Sv.cell(rrow, 1, lab).font = f_base
    for i in range(12):
        v = '=SUMIFS(%s,%s,"%s",%s,%d,%s,1)/1000000' % (c("Стоимость на площадке (макс.), руб"), c("Площадка"), s,
                                                    c("Месяц ввода в консигнацию"), i + 1, c("В лимите (1/0)"))
        cc = Sv.cell(rrow, 2 + i, v); cc.number_format = MLN; cc.border = box; cc.font = f_base
for col in range(5, 14): Sv.column_dimensions[CL(col)].width = 9

# ---------------- анализ вне плана ----------------
Av = wb.create_sheet("Вне плана — анализ", 2)
Av["A1"] = "Позиции, использованные вне плана за три года до года плана"; Av["A1"].font = f_title
Av["A2"] = "Вероятность = 1 − e^(−заказов в год). Минимум у использованных позиций ≈31%: один заказ за 32 месяца."; Av["A2"].font = f_note
hdrA = ["Вероятность использования", "от", "до", "1100: позиций", "1100: нет в плане", "1100: ожидаемо, млн руб/год", "1100: в расчёте",
        "1300: позиций", "1300: нет в плане", "1300: ожидаемо, млн руб/год", "1300: в расчёте"]
for i, h in enumerate(hdrA):
    cc = Av.cell(4, i + 1, h); cc.font = f_head; cc.fill = fill_head; cc.alignment = center
    Av.column_dimensions[CL(i + 1)].width = 16 if i else 26
BK = [("90–100%", 0.9, 1.01), ("70–90%", 0.7, 0.9), ("50–70%", 0.5, 0.7), ("30–50%", 0.3, 0.5), ("10–30%", 0.1, 0.3), ("до 10%", 0.000001, 0.1)]
for i, (lab, lo, hi) in enumerate(BK):
    x = 5 + i
    Av.cell(x, 1, lab).font = f_base
    for col, v in ((2, lo), (3, hi)):
        cc = Av.cell(x, col, v); cc.font = f_in; cc.number_format = PCT0
    for j, s in enumerate(("1100", "1300")):
        pr_ = c("Вероятность использования в году плана")
        cond = '%s,"%s",%s,">="&$B%d,%s,"<"&$C%d' % (c("Площадка"), s, pr_, x, pr_, x)
        vals = ["=COUNTIFS(%s)" % cond, '=COUNTIFS(%s,%s,"Вне плана")' % (cond, c("Источник спроса")),
                "=SUMIFS(%s,%s)/1000000" % (c("Вне плана: ожидаемо, руб/год (без порога)"), cond),
                "=SUMIFS(%s,%s)" % (c("Вне плана в расчёт (1/0)"), cond)]
        for k, v in enumerate(vals):
            cc = Av.cell(x, 4 + j * 4 + k, v); cc.number_format = MLN if k == 2 else INT; cc.border = box; cc.font = f_base
x = 5 + len(BK)
Av.cell(x, 1, "Итого").font = f_bold
for col in range(4, 12):
    cc = Av.cell(x, col, "=SUM(%s5:%s%d)" % (CL(col), CL(col), x - 1)); cc.font = f_bold; cc.fill = fill_sub
    cc.number_format = MLN if col in (6, 10) else INT

# ---------------- код Power Query и инструкция ----------------
mtext = io.open(A.m, encoding="utf-8").read()
sel = ", ".join('"%s"' % s[0] for s in SRC)
ren = ", ".join('{"%s", "%s"}' % (s[0], s[1]) for s in SRC if s[0] != s[1])
typ = ", ".join('{"%s", %s}' % (s[1], {"text": "type text", "num": "type number", "int": "Int64.Type"}[s[2]]) for s in SRC)
mtext = mtext.replace("{{SELECT}}", "{" + sel + "}").replace("{{RENAME}}", "{" + ren + "}").replace("{{TYPES}}", "{" + typ + "}")
io.open(os.path.splitext(A.out)[0] + ".Section1.m", "w", encoding="utf-8").write(mtext)
QUERIES = re.findall(r"^shared\s+([^\s=]+)\s*=", mtext, re.M)

Cd = wb.create_sheet("Код Power Query")
Cd["A1"] = "Код запросов Power Query (для правки или ручного восстановления)"; Cd["A1"].font = f_title
Cd["A2"] = "Каждый блок — отдельный запрос: Данные → Получить данные → Из других источников → Пустой запрос → Расширенный редактор → вставить текст блока; имя запроса — как в заголовке."; Cd["A2"].font = f_note
Cd.column_dimensions["A"].width = 160
blocks = re.split(r"(?m)^(?=shared\s)", mtext.split("\n", 1)[1])
rrow = 4
for b in blocks:
    m_ = re.match(r"shared\s+([^\s=]+)\s*=\s*", b)
    if not m_: continue
    body = b[m_.end():].rstrip().rstrip(";")
    body = re.sub(r"\n\s*//[^\n]*$", "", body)
    cc = Cd.cell(rrow, 1, "Запрос: %s" % m_.group(1)); cc.font = f_head; cc.fill = fill_head; rrow += 1
    for line in body.split("\n"):
        if line.strip().startswith("//"): continue
        Cd.cell(rrow, 1, line).font = Font(name="Consolas", size=9); rrow += 1
    rrow += 1

Hw = wb.create_sheet("Как обновить", 0)
Hw.column_dimensions["A"].width = 140
LINES = [
 ("Модель консигнации CAT в Excel: Power Query + формулы, без Python", f_title), ("", None),
 ("1. Положите выгрузки в одну папку", f_bold),
 ("   • PM-06 из ТОРО: M06_1100_<год>.xlsx и M06_1300_<год>.xlsx — год плана и три года до него (например 2024, 2025, 2026, 2027).", None),
 ("   • MM-M03 остатки: Остатки*.xlsx;  закупка АО «Развитие»: Закупка*.xlsx;  справочник: ЕКМТР*.xlsx;  замены номеров Cat: Замены*.xlsx.", None),
 ("   • Если файлов одного вида несколько, берётся самый свежий по дате изменения.", None),
 ("2. На листе «Параметры» укажите папку, год плана и дату выгрузки закупки (жёлтые ячейки).", f_bold),
 ("3. Один раз разрешите объединять данные книги и файлов:", f_bold),
 ("   Данные → Получить данные → Параметры запроса → Текущая книга → Конфиденциальность → «Игнорировать уровни конфиденциальности». Иначе Excel выдаст ошибку Formula.Firewall.", None),
 ("4. Данные → Обновить все. Первая загрузка — несколько минут: PM-06 — это сотни тысяч строк.", f_bold),
 ("   Запрос «Позиции» перезаливает синие столбцы таблицы на листе «Позиции»; чёрные столбцы — формулы, они растягиваются сами.", None),
 ("5. Результат — лист «Сводка» и «Вне плана — анализ». Параметры расчёта (срок, сервис, порог, лимит…) меняются на листе «Параметры» без обновления запросов.", f_bold),
 ("", None),
 ("Что делает Power Query", f_bold),
 ("   ТОРО_строки — строки PM-06 по технике из таблицы «Техника», только Caterpillar (по изготовителю в заказе или в ЕК МТР), без услуг и нулевых строк.", None),
 ("   План — год плана: количество по месяцам, заказы, поле/цех, CAPEX/OPEX.  ВнеПлана — три прошлых года, виды работ вне ТО и КР.", None),
 ("   Остатки — MM-M03 по заводам из таблицы «Заводы», без складов подрядчиков (категория склада 5), консигнации и АТЗ.", None),
 ("   Закупка — действующие заказы на завод площадки по месяцу поставки; заявки и заказы на согласовании — отдельно.", None),
 ("   Замены — цепочки замен Cat; остаток под взаимозаменяемыми номерами.  Исключено — что отсечено правилами «Исключения» (можно загрузить на лист).", None),
 ("", None),
 ("Если запросы не видны (Данные → Запросы и подключения пусто)", f_bold),
 ("   Создайте их из листа «Код Power Query»: по одному пустому запросу на блок, имя — как в заголовке блока; «Позиции» загрузите в таблицу,", None),
 ("   остальные — «Только создать подключение». Затем скопируйте значения новой таблицы в синие столбцы таблицы «Позиции» — формулы пересчитаются.", None),
 ("", None),
 ("Методика — как в модели v4: план по графику (окно срока пополнения), вне плана — вероятность 1 − e^(−λ) и составной Пуассон;", None),
 ("сначала своё — остаток и поставки площадки без перемещений; стоимость — физический запас на площадке; лимит — по числу заказов на рубль запаса.", None),
]
for i, (t, fnt) in enumerate(LINES):
    Hw.cell(i + 1, 1, t).font = fnt or f_base

wb.calculation.fullCalcOnLoad = True
qsheet_idx = wb.sheetnames.index("Позиции")
if not A.no_pq:
    wb.defined_names["ExternalData_1"] = DefinedName("ExternalData_1", localSheetId=qsheet_idx, hidden=True,
                                                     attr_text="'Позиции'!$A$1:$%s$%d" % (CL(len(SRC)), LASTROW))
tmp = A.out + ".tmp"
wb.save(tmp)

# ---------------- встраивание Power Query ----------------
def mashup(section_text, queries, loaded, fill_cols, fill_count):
    def zipbytes(files):
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
            for n_, data in files: z.writestr(n_, data)
        return bio.getvalue()
    bom = b"\xef\xbb\xbf"
    ct = bom + b'<?xml version="1.0" encoding="utf-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="text/xml" /><Default Extension="m" ContentType="application/x-ms-m" /></Types>'
    pkg = bom + b'<?xml version="1.0" encoding="utf-8"?><Package xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema"><Version>1.0.0.0</Version><MinimumVersion>1.0.0.0</MinimumVersion><Culture>ru-RU</Culture></Package>'
    parts = zipbytes([("[Content_Types].xml", ct), ("Config/Package.xml", pkg), ("Formulas/Section1.m", section_text.encode("utf-8"))])
    perm = bom + b'<?xml version="1.0" encoding="utf-8"?><PermissionList xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema"><CanEvaluateFuturePackages>false</CanEvaluateFuturePackages><FirewallEnabled>false</FirewallEnabled><WorkbookGroupType xsi:nil="true" /></PermissionList>'
    now = datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S.0000000Z")
    items = ['<Item><ItemLocation><ItemType>AllFormulas</ItemType><ItemPath /></ItemLocation><StableEntries /></Item>']
    for qn in queries:
        path = "Section1/" + quote(qn, safe="")
        if qn == loaded:
            ent = [("IsPrivate", "l0"), ("FillEnabled", "l1"), ("FillObjectType", "sTable"), ("FillToDataModelEnabled", "l0"),
                   ("ResultType", "sTable"), ("BufferNextRefresh", "l1"), ("FillTarget", "s" + T), ("FillStatus", "sComplete"),
                   ("FillCount", "l%d" % fill_count), ("FillErrorCode", "sUnknown"), ("FillErrorCount", "l0"),
                   ("FillLastUpdated", "d" + now), ("FillColumnNames", "s" + json.dumps(fill_cols, ensure_ascii=False)),
                   ("AddedToDataModel", "l0"), ("QueryID", "s" + str(uuid.uuid4()))]
        else:
            ent = [("IsPrivate", "l0"), ("FillEnabled", "l0"), ("FillObjectType", "sConnectionOnly"), ("FillToDataModelEnabled", "l0"),
                   ("AddedToDataModel", "l0"), ("QueryID", "s" + str(uuid.uuid4()))]
            if qn.startswith("фн"): ent.append(("ResultType", "sFunction"))
        items.append('<Item><ItemLocation><ItemType>Formula</ItemType><ItemPath>%s</ItemPath></ItemLocation><StableEntries>%s</StableEntries></Item>'
                     % (escape(path), "".join('<Entry Type="%s" Value="%s" />' % (k, escape(v, {'"': "&quot;"})) for k, v in ent)))
    meta_xml = bom + ('<?xml version="1.0" encoding="utf-8"?><LocalPackageMetadataFile xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema"><Items>%s</Items></LocalPackageMetadataFile>' % "".join(items)).encode("utf-8")
    meta_content = zipbytes([("[Content_Types].xml", ct)])
    meta = struct.pack("<I", 0) + struct.pack("<I", len(meta_xml)) + meta_xml + struct.pack("<I", len(meta_content)) + meta_content
    blob = (struct.pack("<I", 0) + struct.pack("<I", len(parts)) + parts + struct.pack("<I", len(perm)) + perm
            + struct.pack("<I", len(meta)) + meta + struct.pack("<I", 0))
    xml = '<?xml version="1.0" encoding="utf-16"?><DataMashup xmlns="http://schemas.microsoft.com/DataMashup">%s</DataMashup>' % base64.b64encode(blob).decode()
    return b"\xff\xfe" + xml.encode("utf-16-le")

def inject(src, dst):
    zin = zipfile.ZipFile(src)
    files = {n_: zin.read(n_) for n_ in zin.namelist()}
    zin.close()
    # таблица «Позиции»
    tname = [n_ for n_ in files if n_.startswith("xl/tables/table") and ('displayName="%s"' % T).encode() in files[n_]][0]
    tid = re.search(rb'<table[^>]* id="(\d+)"', files[tname]).group(1).decode()
    ref = "A1:%s%d" % (CL(NC), LASTROW)
    cols = []
    for j, h in enumerate(ALLCOLS):
        cid = j + 1
        if j < len(SRC):
            cols.append('<tableColumn id="%d" uniqueName="%d" name="%s" queryTableFieldId="%d"/>' % (cid, cid, escape(h, {'"': "&quot;"}), cid))
        else:
            fml = CALC[j - len(SRC)][1]
            cols.append('<tableColumn id="%d" uniqueName="%d" name="%s" queryTableFieldId="%d"><calculatedColumnFormula>%s</calculatedColumnFormula></tableColumn>'
                        % (cid, cid, escape(h, {'"': "&quot;"}), cid, escape(fml)))
    files[tname] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" id="%s" name="%s" displayName="%s" ref="%s" tableType="queryTable" totalsRowShown="0">'
        '<autoFilter ref="%s"/><tableColumns count="%d">%s</tableColumns>'
        '<tableStyleInfo name="TableStyleLight9" showFirstColumn="0" showLastColumn="0" showRowStripes="1" showColumnStripes="0"/></table>'
        % (tid, T, T, ref, ref, NC, "".join(cols))).encode("utf-8")
    trels = tname.replace("xl/tables/", "xl/tables/_rels/") + ".rels"
    files[trels] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/queryTable" Target="../queryTables/queryTable1.xml"/></Relationships>').encode()
    fields = []
    for j, h in enumerate(ALLCOLS):
        if j < len(SRC):
            fields.append('<queryTableField id="%d" name="%s" tableColumnId="%d"/>' % (j + 1, escape(h, {'"': "&quot;"}), j + 1))
        else:
            fields.append('<queryTableField id="%d" dataBound="0" tableColumnId="%d"/>' % (j + 1, j + 1))
    files["xl/queryTables/queryTable1.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="ExternalData_1" connectionId="1" autoFormatId="16" '
        'applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="0">'
        '<queryTableRefresh nextId="%d"><queryTableFields count="%d">%s</queryTableFields></queryTableRefresh></queryTable>'
        % (NC + 1, NC, "".join(fields))).encode("utf-8")
    files["xl/connections.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<connection id="1" keepAlive="1" name="Запрос — %s" description="Соединение с запросом &quot;%s&quot; в книге." type="5" refreshedVersion="8" background="1" saveData="1">'
        '<dbPr connection="Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location=%s;Extended Properties=&quot;&quot;" command="SELECT * FROM [%s]"/>'
        '</connection></connections>' % (T, T, T, T)).encode("utf-8")
    files["customXml/item1.xml"] = mashup(mtext, QUERIES, T, [s[1] for s in SRC], N)
    guid = "{%s}" % str(uuid.uuid4()).upper()
    files["customXml/itemProps1.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        '<ds:datastoreItem ds:itemID="%s" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/></ds:datastoreItem>' % guid).encode()
    files["customXml/_rels/item1.xml.rels"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps1.xml"/></Relationships>').encode()
    wr = files["xl/_rels/workbook.xml.rels"].decode()
    wr = wr.replace("</Relationships>",
        '<Relationship Id="rIdPQ1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/connections" Target="connections.xml"/>'
        '<Relationship Id="rIdPQ2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" Target="../customXml/item1.xml"/></Relationships>')
    files["xl/_rels/workbook.xml.rels"] = wr.encode()
    ctp = files["[Content_Types].xml"].decode()
    ctp = ctp.replace("</Types>",
        '<Override PartName="/xl/connections.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml"/>'
        '<Override PartName="/xl/queryTables/queryTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml"/>'
        '<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/></Types>')
    if 'Extension="xml"' not in ctp:
        ctp = ctp.replace("<Default ", '<Default Extension="xml" ContentType="application/xml"/><Default ', 1)
    files["[Content_Types].xml"] = ctp.encode()
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        order = ["[Content_Types].xml"] + [n_ for n_ in files if n_ != "[Content_Types].xml"]
        for n_ in order: z.writestr(n_, files[n_])

if A.no_pq:
    os.replace(tmp, A.out)
else:
    inject(tmp, A.out); os.remove(tmp)
print("записано", A.out, "| позиций", N, "| столбцов", NC, "| запросов", len(QUERIES))
