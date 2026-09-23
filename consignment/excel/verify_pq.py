# -*- coding: utf-8 -*-
"""Независимая проверка вычисляемых столбцов таблицы «Позиции» (книга должна быть пересчитана)."""
import math, sys
from statistics import NormalDist
from openpyxl import load_workbook
MON = ["янв","фев","мар","апр","май","июн","июл","авг","сен","окт","ноя","дек"]
def pq(p, mu):
    if mu >= 20:
        z = NormalDist().inv_cdf(p); return math.ceil(mu + z*math.sqrt(mu) + (z*z-1)/6 - 0.5)
    k, c, lp = 0, 0.0, -mu
    while True:
        c += math.exp(lp)
        if c >= p: return k
        k += 1; lp += math.log(mu) - math.log(k)
wb = load_workbook(sys.argv[1], data_only=True)
P = wb["Параметры"]; hdr = [c.value for c in P[11]]
par = {}
for rr in (12, 13):
    row = [c.value for c in P[rr]]
    par[str(row[0])] = dict(zip(hdr, row))
ws = wb["Позиции"]; h = [c.value for c in ws[1]]; ix = {k: i for i, k in enumerate(h)}
bad = []; n = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None: break
    n += 1; g = lambda k: row[ix[k]] or 0
    pp = par[str(row[0])]
    L, R, SL, K = pp["Срок пополнения, дн"], pp["Период пересмотра, дн"], pp["Уровень сервиса"], pp["Коэффициент выполнения плана"]
    SWV, THR, HY, KH = pp["Учитывать спрос вне плана (1/0)"], pp["Порог вероятности"], pp["Длительность истории, лет"], pp["Коэффициент к истории"]
    SWR, SWP, BUD = pp["Учитывать остаток «Развития» (1/0)"], pp["Учитывать замены Cat (1/0)"], pp["Лимит стоимости, руб (0 — без лимита)"]
    SWD, SWQ, DLY = pp["Учитывать поставки по заказам (1/0)"], pp["Учитывать неподтверждённые (1/0)"], pp["Задержка поставок, мес"]
    W = max(1, min(12, round(L / 30.4)))
    ho = sum(g("Вне плана: заказов (год −%d)" % k) for k in (3, 2, 1)); hq = sum(g("Вне плана: шт (год −%d)" % k) for k in (3, 2, 1))
    lam = ho / HY * KH; qb = hq / ho if ho else 0; P_ = 1 - math.exp(-lam)
    inc = 1 if (SWV == 1 and lam > 0 and P_ >= THR) else 0
    dv = inc * max(0, lam * qb - g("в т.ч. внеплановые виды работ в плане, шт") * K)
    dp = g("План, шт") * K; dem = dp + dv
    m = [g("План %s, шт" % x) * K for x in MON]
    mplan = max(sum(m[i:i + W]) for i in range(12))
    qe = dv / lam if (dv > 0 and lam > 0) else 0
    muL = lam * L / 365 if dv > 0 else 0; muLR = lam * (L + R) / 365 if dv > 0 else 0
    mn = 0 if dem <= 0 else math.ceil(mplan + (pq(SL, muL) * qe if muL > 0 else 0) - 1e-6)
    mp = math.ceil(max(g("Макс. разовый отбор в плане, шт"), g("Вне плана: макс. разовый отбор, шт") if dv > 0 else 0))
    mx = 0 if dem <= 0 else max(math.ceil(mplan + dp * R / 365 + (pq(SL, muLR) * qe if muLR > 0 else 0) - 1e-6), mn + 1, mp)
    own = g("Остаток БЕ площадки, шт") + g("Остаток АО «Развитие» без подрядчиков, шт") * SWR + g("Остаток по заменам Cat, шт") * SWP
    dm = [g("Поставка %s, шт" % x) for x in MON]; sup = []
    for k in range(1, 13):
        s = own + SWD * ((g("Поставки: до года плана и просроченные, шт") if k >= DLY else 0) + (sum(dm[:k - DLY]) if k - DLY >= 1 else 0))
        s += SWQ * (g("Неподтверждённые: заявки и заказы на согласовании, шт") if k >= min(12, W + 1 + DLY) else 0)
        sup.append(s)
    mon, cum = 13, 0
    if dem > 0:
        for i in range(12):
            cum += m[i] + dv / 12
            if cum > sup[i] - mn + 1e-9: mon = i + 1; break
    edl = dem * L / 365
    on = 0 if dem <= 0 else max(math.ceil(mx - edl - 1e-6), mp, 1)
    exp = {"Вероятность использования в году плана": P_, "Min, шт": mn, "Max, шт": mx, "Месяц ввода в консигнацию": mon,
           "Макс. запас на площадке, шт": on, "Своё предложение на год (остаток + поставки), шт": sup[11]}
    for k, v in exp.items():
        got = row[ix[k]]
        if got is None or abs(got - v) > 1e-6 * max(1, abs(v)):
            bad.append((row[0], row[2], k, got, v)); break
print("проверено позиций:", n, "| расхождений:", len(bad))
for b in bad[:8]: print("   ", b)
sys.exit(1 if bad else 0)
