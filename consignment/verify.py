# -*- coding: utf-8 -*-
"""Независимая проверка книги консигнации: пересчитывает ключевые столбцы на Python
и сравнивает с формулами Excel.

    python consignment/verify.py <книга.xlsx>

Книга должна быть пересчитана (LibreOffice или Excel) — openpyxl читает кэш значений.
Проверяются: вероятность вне плана, Min, Max, запас на площадке (макс. и средний),
своё предложение на 2027 и месяц ввода в консигнацию. Код выхода 1 — есть расхождения.
"""
import math, sys
from statistics import NormalDist
from openpyxl import load_workbook

MON = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]

def pq(p, mu):
    """Квантиль Пуассона так же, как в книге: точно при μ<20, Корниш–Фишер при μ≥20."""
    if mu >= 20:
        z = NormalDist().inv_cdf(p)
        return math.ceil(mu + z * math.sqrt(mu) + (z * z - 1) / 6 - 0.5)
    k, c, lp = 0, 0.0, -mu
    while True:
        c += math.exp(lp)
        if c >= p: return k
        k += 1; lp += math.log(mu) - math.log(k)

def check(ws):
    hdr = [c.value for c in ws[6]]; ix = {h: i for i, h in enumerate(hdr)}
    par = {ws.cell(2, c).value: ws.cell(3, c).value for c in range(2, 22) if ws.cell(2, c).value}
    L, R, SL, K, THR, HY, KH = (par[k] for k in ("L", "R", "SL", "K", "THR", "HY", "KH"))
    SWV, SWR, SWP, SWD, SWQ, DLY, W = (par[k] for k in ("SWV", "SWR", "SWP", "SWD", "SWQ", "DLY", "W, мес"))
    bad, n = [], 0
    for row in ws.iter_rows(min_row=7, values_only=True):
        if row[0] is None: break
        n += 1
        g = lambda h: row[ix[h]] or 0
        ho = sum(g("Вне плана: заказов %s" % y) for y in ("2024", "2025", "2026 (янв–авг)"))
        hq = sum(g("Вне плана: шт %s" % y) for y in ("2024", "2025", "2026 (янв–авг)"))
        lam = ho / HY * KH; qb = hq / ho if ho else 0; P = 1 - math.exp(-lam)
        inc = 1 if (SWV == 1 and lam > 0 and P >= THR) else 0
        dv = inc * max(0, lam * qb - g("в т.ч. внеплановые виды работ в плане, шт") * K)
        dp = g("План 2027, шт") * K; dem = dp + dv
        m = [g("План %s, шт" % x) * K for x in MON]
        mplan = max(sum(m[i:i + W]) for i in range(12))
        qe = dv / lam if (dv > 0 and lam > 0) else 0
        muL = lam * L / 365 if dv > 0 else 0; muLR = lam * (L + R) / 365 if dv > 0 else 0
        mn = 0 if dem <= 0 else math.ceil(mplan + (pq(SL, muL) * qe if muL > 0 else 0) - 1e-6)
        mp = max(g("Макс. разовый отбор в плане, шт"), g("Вне плана: макс. разовый отбор, шт") if dv > 0 else 0)
        mx = 0 if dem <= 0 else max(math.ceil(mplan + dp * R / 365 + (pq(SL, muLR) * qe if muLR > 0 else 0) - 1e-6),
                                    mn + 1, math.ceil(mp))
        own = g("Остаток БЕ площадки, шт") + g("Остаток АО «Развитие» без подрядчиков, шт") * SWR + g("Остаток по заменам Cat, шт") * SWP
        dm = [g("Поставка %s 2027, шт" % x) for x in MON]
        sup = []
        for mm in range(1, 13):
            s = own + SWD * ((g("Поставки: до 2027 и просроченные, шт") if mm >= DLY else 0) + (sum(dm[:mm - DLY]) if mm - DLY >= 1 else 0))
            s += SWQ * (g("Неподтверждённые: заявки и заказы на согласовании, шт") if mm >= min(12, W + 1 + DLY) else 0)
            sup.append(s)
        mon, cum = 13, 0
        if dem > 0:
            for i in range(12):
                cum += m[i] + dv / 12
                if cum > sup[i] - mn + 1e-9: mon = i + 1; break
        edl = dem * L / 365
        on = 0 if dem <= 0 else max(math.ceil(mx - edl - 1e-6), math.ceil(mp), 1)
        inn = 1 if (dem > 0 and mon <= 12) else 0
        avg = inn * min(on, max(0, mn - edl) + (mx - mn) / 2) * g("Цена, руб/шт")
        exp = {"Вероятность использования в 2027": round(P, 6), "Min, шт": mn, "Max, шт": mx,
               "Месяц ввода в консигнацию": mon, "Макс. запас на площадке, шт": on,
               "Своё предложение на 2027 (остаток + поставки), шт": round(sup[11], 6),
               "Средний запас на площадке, руб": round(avg, 2)}
        for h, v in exp.items():
            got = row[ix[h]]
            got = round(got, 6 if isinstance(v, float) and h != "Средний запас на площадке, руб" else 2) if isinstance(got, float) else got
            if (abs(got - v) > 0.01) if isinstance(v, float) else got != v:
                bad.append((row[1], h, got, v)); break
    return n, bad

def main():
    wb = load_workbook(sys.argv[1], data_only=True)
    total_bad = 0
    for t in ("Красноярск", "Алдан"):
        n, bad = check(wb[t])
        total_bad += len(bad)
        print("%s: проверено %d позиций, расхождений %d" % (t, n, len(bad)))
        for b in bad[:5]: print("   ", b)
    sys.exit(1 if total_bad else 0)

if __name__ == "__main__":
    main()
